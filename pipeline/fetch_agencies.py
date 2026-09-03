"""Measures Statistics Sweden does not publish: the National Council for Crime
Prevention, SALAR and the National Board of Health and Welfare (via Kolada/RKA),
Transport Analysis and the Agency for Growth Policy Analysis.

Part of mandatperioden - https://github.com/klahr/mandatperioden
Copyright (C) 2026 Joachim Klahr

This program is free software: you can redistribute it and/or modify it
under the terms of the GNU General Public License as published by the Free
Software Foundation, either version 3 of the License, or (at your option)
any later version.

This program is distributed in the hope that it will be useful, but
WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the GNU General
Public License for more details. You should have received a copy of the
License along with this program; see the file LICENSE or
<https://www.gnu.org/licenses/>.
"""

import json, os, time, hashlib, urllib.request, urllib.error, warnings
from collections import defaultdict
import openpyxl
warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "cache_alt")
XL = os.path.join(HERE, "sources", "bra")
ATERFALL_URL = {
    "aterfall_1ar_2014_2023": "https://bra.se/download/18.5d0a8fbf19ddcaf2a6c2173f/1778480930401/p811La-2014-2023.xlsx",
    "aterfall_1ar_2013_2022": "https://bra.se/download/18.388a7da196d6c2b1c49b21/1747660630514/p811La-2013-2022.xlsx",
    "aterfall_3ar_2009_2018": "https://bra.se/download/18.1bcc29ae199371e6e753026e/1758634024334/S811La-2009-2018.xlsx"}
TRAFA_URL = ("https://www.trafa.se/globalassets/statistik/bantrafik/"
             "punktlighet-pa-jarnvag/2026/punktlighet-pa-jarnvag-2025.xlsx")
BRA_URL = {
 "uppklarade.xlsx": "https://bra.se/download/18.4450dd8019ed44e114546eae/1783421446755/10La_uppklbr_10_ar.xlsx",
 "dodligt_vald.xlsx": "https://bra.se/download/18.5b3bbb9a19d24bdce4b157ff/1774872877489/Tabell%2020_2002-2025.xlsx",
 "dv_skjut.xlsx": "https://bra.se/download/18.5b3bbb9a19d24bdce4b15801/1774872877607/Tabell%2022_2016-2025.xlsx",
 "hatbrott.xlsx": "https://bra.se/download/18.125e930a19b6e6f26ff4e23e/1783940787861/Tabellsamling%20Polisanm%C3%A4lda%20hatbrott%202020-2024.xlsx",
}
os.makedirs(CACHE, exist_ok=True)

LOG = []
def _log(method, url, owner, note=None, key=None):
    LOG.append({"method": method, "url": url, "owner": owner, "note": note, "key": key})

def fetch(url):
    fn = os.path.join(CACHE, hashlib.sha256(url.encode()).hexdigest()[:24] + ".json")
    if os.path.exists(fn):
        return json.load(open(fn, encoding="utf-8"))
    for a in range(4):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "mandate-stats/1.0"})
            d = json.load(urllib.request.urlopen(req, timeout=120)); break
        except Exception:
            if a == 3: raise
            time.sleep(3*(a+1))
    json.dump(d, open(fn, "w", encoding="utf-8"), ensure_ascii=False); time.sleep(0.3)
    return d

KOL_META = {}
def kolada(kpi):
    _log("GET", f"https://api.kolada.se/v3/kpi/{kpi}", "kolada",
         "nyckeltalets definition och statistikansvarig", kpi)
    _log("GET", f"https://api.kolada.se/v3/data/kpi/{kpi}/municipality/0000", "kolada",
         "hela tidsserien för riket (municipality 0000)", kpi)
    KOL_META[kpi] = fetch(f"https://api.kolada.se/v3/kpi/{kpi}")["values"][0]
    d = fetch(f"https://api.kolada.se/v3/data/kpi/{kpi}/municipality/0000")
    out = {}
    for x in d.get("values", []):
        v = x["values"][0]["value"] if x.get("values") else None
        # Series around 1.9 hospital beds lose the entire change at one decimal.
        if v is not None: out[str(x["period"])] = round(v, 4)
    return out

def sos(diagnos, first=2008):
    u = ("https://sdb.socialstyrelsen.se/api/v1/sv/dodsorsaker/resultat"
         f"?diagnos={diagnos}&region=0&kon=3&matt=1")
    _log("GET", u, "sos", "riket, båda könen, antal döda – summeras över åldersgrupper", diagnos)
    d = fetch(u)
    tot, seen = defaultdict(float), set()
    for r in d["data"]:
        if r["regionId"] != 0 or r["konId"] != 3 or r["mattId"] != 1: continue
        try:
            tot[str(r["ar"])] += float(str(r["varde"]).replace(",", ".")); seen.add(str(r["ar"]))
        except ValueError:
            pass
    return {k: int(v) for k, v in sorted(tot.items()) if k in seen and int(k) >= first}

def bra_rate(sheet, brottstyp):
    ws = openpyxl.load_workbook((_log("GET", BRA_URL["uppklarade.xlsx"], "bra", "bladet Statistik personuppklaringsproc respektive lagföringsprocent", "uppklarade.xlsx") or os.path.join(XL, "uppklarade.xlsx")), data_only=True)[sheet]
    years = [str(c.value) for c in ws[2][2:] if c.value]
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[1] and str(row[1]).strip() == brottstyp:
            out = {}
            for y, v in zip(years, row[2:]):
                try: out[y] = round(float(v), 2)
                except (TypeError, ValueError): pass
            return out
    return {}

def bra_dodligt():
    ws = openpyxl.load_workbook((_log("GET", BRA_URL["dodligt_vald.xlsx"], "bra", "bladet Statistik, tabell 20", "dodligt_vald.xlsx") or os.path.join(XL, "dodligt_vald.xlsx")), data_only=True)["Statistik"]
    ant, per = {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        try: y = str(int(row[0]))
        except (TypeError, ValueError): continue
        if int(y) < 2008: continue
        if row[1] is not None: ant[y] = int(row[1])
        if row[2] is not None: per[y] = round(float(row[2]), 2)
    return ant, per

def bra_skjutvapen():
    wb = openpyxl.load_workbook((_log("GET", BRA_URL["dv_skjut.xlsx"], "bra", "första bladet, tabell 22", "dv_skjut.xlsx") or os.path.join(XL, "dv_skjut.xlsx")), data_only=True)
    ws = wb[wb.sheetnames[0]]
    tot, med, utan = {}, {}, {}
    def iv(x):
        try: return int(x)
        except (TypeError, ValueError): return None
    for row in ws.iter_rows(min_row=2, values_only=True):
        try: y = str(int(row[0]))
        except (TypeError, ValueError): continue
        if int(y) < 2008: continue
        if iv(row[1]) is not None: tot[y] = iv(row[1])
        if iv(row[2]) is not None: med[y] = iv(row[2])
        if iv(row[5]) is not None: utan[y] = iv(row[5])
    return tot, med, utan

def bra_aterfall(fil, matt):
    """Sheet "Statistik": sex in column A, measure in column B, one year per
    column thereafter."""
    _log("GET", ATERFALL_URL[fil], fil, f"Samtliga, {matt}", f"{fil}.xlsx")
    ws = openpyxl.load_workbook(os.path.join(HERE, "sources", "xlsx", fil + ".xlsx"),
                                data_only=True)["Statistik"]
    rows = list(ws.iter_rows(min_row=1, max_row=30, values_only=True))
    head = next(r for r in rows if r[0] == "Kön")
    years = {i: str(int(v)) for i, v in enumerate(head) if isinstance(v, (int, float, str))
             and str(v).strip().isdigit() and len(str(v).strip()) == 4}
    row = next(r for r in rows if r[0] == "Samtliga" and r[1] == matt)
    return {y: round(float(row[i]), 3) for i, y in years.items()
            if isinstance(row[i], (int, float))}

def trafa_punktlighet():
    """Column K = share arriving within 5 minutes (RT+5), the official measure.
    Figures before 2013 are not quality-assured and are left out."""
    _log("GET", TRAFA_URL, "trafa", "tabell A, månadsvis punktlighet RT+5", "punktlighet.xlsx")
    ws = openpyxl.load_workbook(os.path.join(HERE, "sources", "xlsx", "punktlighet.xlsx"),
                                data_only=True)["Tabell A"]
    MON = {"jan":1,"feb":2,"mar":3,"apr":4,"maj":5,"jun":6,
           "jul":7,"aug":8,"sep":9,"okt":10,"nov":11,"dec":12}
    out = {}
    for row in ws.iter_rows(min_row=9, values_only=True):
        y, mo, v = row[1], row[2], row[10]
        if not isinstance(y, (int, float)) or not isinstance(v, (int, float)): continue
        if int(y) < 2013: continue
        m = MON.get(str(mo).strip().lower()[:3])
        if m: out[f"{int(y)}M{m:02d}"] = round(float(v), 2)
    return out

def bra_hatbrott():
    wb = openpyxl.load_workbook((_log("GET", BRA_URL["hatbrott.xlsx"], "bra", "bladen 4A och 5A", "hatbrott.xlsx") or os.path.join(XL, "hatbrott.xlsx")), data_only=True)
    ws = wb["4A"]
    mark = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        try: y = str(int(row[0]))
        except (TypeError, ValueError): continue
        v = str(row[1]).replace(" ", "").replace(" ", "").replace("*", "")
        try: mark[y] = int(v)
        except ValueError: pass
    ws = wb["5A"]
    yrs = [str(int(v)) for v in [c.value for c in ws[3]] if isinstance(v, (int, float))]
    motiv = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] and str(row[0]).strip() == "Totalt":
            vals = [v for v in row[1:] if isinstance(v, (int, float))]
            motiv = {y: int(vals[i*2]) for i, y in enumerate(yrs) if i*2 < len(vals)}
    return mark, motiv

POP_Y = json.load(open(os.path.join(HERE, "data", "data.json"), encoding="utf-8"))["pop"]
def per100k(ser, nd=2):
    return {k: round(v / POP_Y[k[:4]] * 1e5, nd) for k, v in ser.items()
            if v is not None and k[:4] in POP_Y}

IND, SRC = [], {}
_seen = 0
def src(sid, **kw): SRC[sid] = kw
def add(**kw):
    global _seen
    kw["reqs"] = LOG[_seen:]
    _seen = len(LOG)
    IND.append(kw)

src("SKR-VOD", namn="Väntetider i vården – nationella väntetidsdatabasen",
    producent="Sveriges Kommuner och Regioner (SKR)",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), Rådet för främjande av kommunala analysers öppna databas.",
    url="https://www.vantetider.se/")
src("BRA-HANDL", namn="Handlagda brott – tabell 10La, uppklarade brott tio år",
    producent="Brottsförebyggande rådet (Brå)",
    kanal="Nedladdad tabellfil (xlsx) från Brås statistiksida för handlagda brott.",
    url="https://bra.se/statistik/kriminalstatistik/handlagda-brott.html")
src("BRA-DV", namn="Konstaterade fall av dödligt våld – tabell 20 och 22",
    producent="Brottsförebyggande rådet (Brå)",
    kanal="Nedladdade tabellfiler (xlsx) från Brås statistiksida för konstaterade fall av dödligt våld.",
    url="https://bra.se/statistik/kriminalstatistik/konstaterade-fall-av-dodligt-vald.html")
src("BRA-HAT", namn="Polisanmälda hatbrott – tabellsamling 2020–2024",
    producent="Brottsförebyggande rådet (Brå)",
    kanal="Nedladdad tabellsamling (xlsx) från Brås statistiksida för hatbrottsstatistik.",
    url="https://bra.se/statistik/statistiska-undersokningar/hatbrottsstatistik.html")
src("BRA-KOLADA", namn="Anmälda brott och Nationella trygghetsundersökningen (NTU)",
    producent="Brottsförebyggande rådet (Brå)",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), som publicerar Brås kommun- och riksstatistik.",
    url="https://bra.se/statistik.html")
src("SKR-VERKS", namn="Regionernas verksamhetsstatistik – disponibla vårdplatser",
    producent="Sveriges Kommuner och Regioner (SKR)",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), riket.",
    url="https://www.kolada.se/")
src("SKOLVERKET", namn="Grundskolans betygsstatistik – meritvärde årskurs 9",
    producent="Skolverket",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), riket.",
    url="https://www.skolverket.se/skolutveckling/statistik")
src("KRONOFOGDEN", namn="Skuldsatta fysiska personer hos Kronofogden",
    producent="Kronofogden",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), riket.",
    url="https://kronofogden.se/om-kronofogden/statistik")
src("BRA-ATERFALL", namn="Återfall i brott – tabell s811 (slutlig) och p811 (preliminär)",
    producent="Brottsförebyggande rådet (Brå)",
    kanal="Nedladdade tabellfiler (xlsx) från Brås urvalsverktyg för återfallsstatistik.",
    url="https://bra.se/statistik/statistik-om-rattsvasendet/aterfall-i-brott")
src("TRAFA", namn="Punktlighet på järnväg – tabell A, månadsvis",
    producent="Trafikanalys",
    kanal="Nedladdad tabellfil (xlsx) från Trafikanalys statistiksida. Grunddata kommer från Trafikverkets uppföljningssystem för tågtrafik.",
    url="https://www.trafa.se/bantrafik/punktlighet-pa-jarnvag/")
src("TVA-KONKURS", namn="Företagskonkurser",
    producent="Tillväxtanalys",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), riket.",
    url="https://tillvaxtanalys.se/statistik/konkurser.html")
src("SOS-ALDRE", namn="Statistik om socialtjänstinsatser till äldre",
    producent="Socialstyrelsen",
    kanal="Hämtat maskinellt ur Kolada (api.kolada.se/v3), riket.",
    url="https://www.socialstyrelsen.se/statistik-och-data/statistik/statistikamnen/socialtjanstinsatser-till-aldre/")
src("SOS-DOD", namn="Dödsorsaksregistret – statistikdatabasen för dödsorsaker",
    producent="Socialstyrelsen",
    kanal="Hämtat maskinellt ur Socialstyrelsens öppna statistik-API (sdb.socialstyrelsen.se/api/v1/sv/dodsorsaker).",
    url="https://sdb.socialstyrelsen.se/")

add(key="vardkoer", source="SKR", tables=["SKR-VOD"],
    kpi=["N79223", "N79221", "N79224", "N79222", "U79049", "U72552"],
    series=kolada("N79223"),
    extra={"Väntande högst 90 dagar till första kontakt i specialiserad vård, procent": kolada("N79221"),
           "Genomförd operation/åtgärd inom 90 dagar, procent": kolada("N79224"),
           "Genomförd första kontakt inom 90 dagar, procent": kolada("N79222"),
           "Genomförda första besök inom 90 dagar, allmänpsykiatri, procent (lång serie)": kolada("U79049"),
           "Genomförda första besök inom 90 dagar, barn- och ungdomspsykiatri, procent (lång serie)": kolada("U72552")})
add(key="vantetider", source="SKR", tables=["SKR-VOD"], kpi=["N79242", "N79243", "N79241"],
    series=kolada("N79242"),
    extra={"Väntetid för genomförd operation/åtgärd, mediandagar": kolada("N79243"),
           "Väntetid för genomförd första kontakt i specialiserad vård, mediandagar": kolada("N79241")})

pers_mord = bra_rate("Statistik personuppklaringsproc", "Fullbordat mord och dråp, totalt")
add(key="uppklarade_brott", source="BRÅ", tables=["BRA-HANDL"],
    series=bra_rate("Statistik personuppklaringsproc", "SAMTLIGA BROTT"),
    extra={"Lagföringsprocent, samtliga brott": bra_rate("Statistik lagföringsprocent", "SAMTLIGA BROTT"),
           "Personuppklaringsprocent, brott mot person":
               bra_rate("Statistik personuppklaringsproc", "Brott mot person, totalt"),
           "Personuppklaringsprocent, fullbordat mord och dråp": pers_mord})

dv_ant, dv_per = bra_dodligt()
# Brå's own per-100,000 rate has one decimal, so 2018 and 2022 both read 1.1.
# Computed here from Brå's counts and SCB's population; Brå's figure is kept as
# a control series.
add(key="mord", source="BRÅ", tables=["BRA-DV", "SOS-DOD", "BRA-HANDL"], series=per100k(dv_ant),
    extra={"Konstaterade fall, antal (utan befolkningsjustering)": dv_ant,
           "Brås eget publicerade tal per 100 000 inv (en decimal)": dv_per,
           "Döda av övergrepp av annan person (ICD-10 X85–Y09), per 100 000 inv": per100k(sos("2027")),
           "Döda av övergrepp av annan person, antal": sos("2027"),
           "Personuppklaringsprocent, fullbordat mord och dråp": pers_mord})
dv_tot, dv_med, dv_utan = bra_skjutvapen()
add(key="skjutvapenvald", source="BRÅ", tables=["BRA-DV"], series=per100k(dv_med),
    extra={"Konstaterade fall med skjutvapen, antal (utan befolkningsjustering)": dv_med,
           "Konstaterade fall utan skjutvapen, per 100 000 inv": per100k(dv_utan),
           "Samtliga konstaterade fall, per 100 000 inv": per100k(dv_tot)})

hat_mark, hat_motiv = bra_hatbrott()
add(key="hatbrott", source="BRÅ", tables=["BRA-HAT"], series=per100k(hat_mark, 1),
    extra={"Hatbrottsmarkerade anmälningar, antal (utan befolkningsjustering)": hat_mark,
           "Brottsanmälningar med identifierade hatbrottsmotiv enligt Brås granskning, per 100 000 inv":
               per100k(hat_motiv, 1),
           "Brottsanmälningar med identifierade hatbrottsmotiv, antal": hat_motiv})

add(key="anmalda_brott", source="BRÅ", tables=["BRA-KOLADA"],
    kpi=["N07540", "N07403", "N07546", "N07544"], series={k: round(v) for k, v in kolada("N07540").items()},
    extra={"Anmälda våldsbrott, per 100 000 inv": kolada("N07403"),
           "Anmälda bostadsinbrott, per 100 000 inv": kolada("N07546"),
           "Anmälda narkotikabrott, per 100 000 inv": kolada("N07544")})
add(key="ntu_otrygghet", source="BRÅ", tables=["BRA-KOLADA"],
    kpi=["N07611", "N07620", "N07605", "N07606", "N07633"], series=kolada("N07611"),
    extra={"Avstått från någon aktivitet pga oro för brott, procent": kolada("N07620"),
           "Självrapporterad utsatthet för brott mot enskild person, procent": kolada("N07633"),
           "Förtroende för att polisen behandlar brottsutsatta bra, procent": kolada("N07606"),
           "Förtroende för att rättsväsendet hanterar brottsmisstänkta rättvist, procent": kolada("N07605")})

add(key="vardplatser", source="SKR", tables=["SKR-VERKS"], kpi=["N70845", "N72816", "N71816"],
    series=kolada("N70845"),
    extra={"Disponibla vårdplatser i specialiserad somatisk vård, per 1 000 inv": kolada("N72816"),
           "Disponibla vårdplatser i sluten primärvård, per 1 000 inv": kolada("N71816")})
add(key="meritvarde", source="SKOLVERKET", tables=["SKOLVERKET"], kpi=["N15507", "N15566"],
    series=kolada("N15507"),
    extra={"Meritvärde exkl. nyinvandrade och elever med okänd bakgrund": kolada("N15566")})
_sk = kolada("N00989")
add(key="skuldsatta", source="KRONOFOGDEN", tables=["KRONOFOGDEN"], kpi=["N00989", "N00990"],
    series={k: round(v * 10, 2) for k, v in _sk.items()},
    extra={"Andel skuldsatta, procent": _sk,
           "Medianskuld hos Kronofogden, kr": kolada("N00990")})
add(key="medianskuld", source="KRONOFOGDEN", tables=["KRONOFOGDEN"], kpi=["N00990", "N00989"],
    series=kolada("N00990"),
    extra={"Skuldsatta invånare 18+ hos Kronofogden, procent": kolada("N00989")})

# Two vintages with different adjudication follow-up: final statistics for
# index years 2009-2018, preliminary for 2019-2023. The preliminary run about a
# percentage point lower, hence separate segments. Brå notes that the 2018 drop
# in the preliminary series is absent from the final one, so preliminary 2018 is
# not used as a baseline.
_slutlig = bra_aterfall("aterfall_3ar_2009_2018", "Andel återfall, 1 år")
_prel = {**bra_aterfall("aterfall_1ar_2013_2022", "Andel återfall, 1 år"),
         **bra_aterfall("aterfall_1ar_2014_2023", "Andel återfall, 1 år")}
_prel_ny = {y: v for y, v in _prel.items() if int(y) >= 2019}
add(key="aterfall", source="BRÅ", tables=["BRA-ATERFALL"],
    series={**_slutlig, **_prel_ny},
    segments=[sorted(_slutlig), sorted(_prel_ny)],
    extra={"Andel återfall inom tre år, slutlig statistik, procent":
               bra_aterfall("aterfall_3ar_2009_2018", "Andel återfall, 3 år"),
           "Andel återfall inom ett år, preliminär statistik, procent": _prel,
           "Personer med ingångshändelse, slutlig statistik, antal":
               bra_aterfall("aterfall_3ar_2009_2018", "Antal personer med ingångshändelse")})

add(key="punktlighet", source="TRAFA", tables=["TRAFA"], series=trafa_punktlighet(), extra={})
add(key="konkurser", source="TVA", tables=["TVA-KONKURS"], kpi=["N00926", "N01006"],
    series=kolada("N00926"),
    extra={"Anställda i konkursade företag, per 1 000 inv 16–64 år": kolada("N01006")})
add(key="aldreomsorg", source="SOS", tables=["SOS-ALDRE"], kpi=["N20892", "N20891", "N23805"],
    series=kolada("N20892"),
    extra={"Invånare 65+ i särskilt boende eller med hemtjänst, procent": kolada("N20891"),
           "Brukare 65+ i särskilt boende, antal": kolada("N23805")})

nark, alko = sos("Spec1"), sos("Alk")
add(key="narkotikadodlighet", source="SOS", tables=["SOS-DOD"], series=per100k(nark),
    extra={"Antal döda (utan befolkningsjustering)": nark,
           "Alkoholrelaterad dödlighet, per 100 000 inv": per100k(alko),
           "Alkoholrelaterad dödlighet, antal döda": alko})
alko = sos("Alk")
add(key="alkoholdodlighet", source="SOS", tables=["SOS-DOD"], series=per100k(alko),
    extra={"Antal döda (utan befolkningsjustering)": alko,
           "Narkotika- och läkemedelsförgiftningar, per 100 000 inv": per100k(nark),
           "Narkotika- och läkemedelsförgiftningar, antal döda": nark})

json.dump({"requests": LOG}, open(os.path.join(HERE, "data", "requests_alt.json"),
                                  "w", encoding="utf-8"), ensure_ascii=False, indent=1)
json.dump({"indicators": IND, "sources": SRC, "kolada_meta": KOL_META},
          open(os.path.join(HERE, "data", "alt.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("alt-mätpunkter:", len(IND))
for i in IND:
    ks = sorted(i["series"])
    print(f'  {i["key"]:22s} [{i["source"]}] n={len(ks):3d}  {ks[0] if ks else "-"} .. {ks[-1] if ks else "-"}')
